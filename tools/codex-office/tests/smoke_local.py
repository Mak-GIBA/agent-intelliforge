"""Opt-in real local smoke: creates synthetic files, never uses user documents."""
from pathlib import Path
import json
import sys
import traceback

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from common import OfficeError, file_sha, json_bytes, new_file, no_symlinks
import office


def smoke(out: Path, cfg: dict):
    no_symlinks(out)
    if out.exists():raise OfficeError('Smoke output directory must not already exist')
    out.mkdir(parents=True)
    from PIL import Image,ImageDraw
    from docx import Document
    from docx.shared import Inches,Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from pptx import Presentation
    from pptx.util import Inches as PI, Pt as PP
    from pypdf import PdfWriter

    image=out/'fixture.png'
    pic=Image.new('RGB',(300,150),'white');draw=ImageDraw.Draw(pic)
    draw.rectangle((10,10,120,140),outline='black',width=4)
    draw.line((140,120,270,30),fill='black',width=5)
    pic.save(image)

    word=out/'日本語 report original.docx'
    doc=Document();normal=doc.styles['Normal'];normal.font.name='Noto Sans CJK JP';normal.font.size=Pt(11)
    for name in ['Normal','Heading 1']:
        style=doc.styles[name]
        rpr=style.element.get_or_add_rPr();fonts=rpr.find(qn('w:rFonts'))
        if fonts is None:fonts=OxmlElement('w:rFonts');rpr.append(fonts)
        fonts.set(qn('w:eastAsia'),'Noto Sans CJK JP')
    doc.add_heading('日本語の報告書',0)
    doc.add_paragraph('DRAFT')
    doc.add_paragraph('原図を保ち、必要な文章だけを修正します。')
    table=doc.add_table(rows=1, cols=2);table.style='Table Grid'
    table.rows[0].cells[0].text='確認項目';table.rows[0].cells[1].text='状態'
    row=table.add_row();row.cells[0].text='日本語の表示';row.cells[1].text='要確認'
    doc.add_picture(str(image),width=Inches(3));doc.save(word)

    deck=out/'日本語 slides original.pptx'
    prs=Presentation();prs.slide_width=PI(10);prs.slide_height=PI(5.625)
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    box=slide.shapes.add_textbox(PI(.6),PI(.3),PI(8.8),PI(.7))
    run=box.text_frame.paragraphs[0].add_run();run.text='日本語のスライド';run.font.name='Noto Sans CJK JP';run.font.size=PP(28)
    box=slide.shapes.add_textbox(PI(.6),PI(1.3),PI(8.8),PI(.6))
    run=box.text_frame.paragraphs[0].add_run();run.text='DRAFT';run.font.name='Arial';run.font.size=PP(22)
    slide.shapes.add_picture(str(image),PI(.6),PI(2.2),width=PI(4))
    prs.save(deck)

    pdf=out/'sample.pdf';w=PdfWriter();w.add_blank_page(300,200);w.add_blank_page(320,200);w.write(pdf);w.close()
    results=[]
    for src in [word,deck]:
        original_hash=file_sha(src);copy=src.with_name('edited'+src.suffix)
        office.prepare(src,copy)
        tool_result={'kind':'officecli_edit','source':src.name,'status':'NOT_RUN'}
        if cfg.get('officecli',{}).get('path'):
            try:
                office.office_command(cfg,['set',str(copy),'/','--find','DRAFT','--replace','FINAL','--json'])
                office.office_command(cfg,['close',str(copy)])
                validation=office.office_command(cfg,['validate',str(copy)])
                readback=office.office_command(cfg,['view',str(copy),'text'])
                if 'FINAL' not in readback.stdout or 'DRAFT' in readback.stdout:raise OfficeError('Text replacement failed read-back')
                diff=office.compare(src,copy,True)
                if diff['issues']:raise OfficeError('Media or structure preservation failed')
                tool_result.update(status='PASS',validation=validation.stdout[:2000],changed_parts=diff.get('changed_parts'))
            except Exception as e:tool_result.update(status='FAIL',error=str(e))
        else:
            tool_result['reason']='No installed OfficeCLI; not substituted by a mock in this smoke.'
        if file_sha(src)!=original_hash:raise OfficeError('Original fixture was modified')
        results.append(tool_result)
        try:
            r=office.render(copy,out/(src.suffix[1:]+'-pages'),None,120,'auto')
            results.append({'kind':'real_render','source':copy.name,'status':'PASS','pages':r['total_pages'],'visual_review':'NOT_PERFORMED'})
        except Exception as e:results.append({'kind':'real_render','source':copy.name,'status':'FAIL','error':str(e)})
    try:
        r=office.render(pdf,out/'pdf-pages',None,120,'auto')
        office.pdf_edit('select',[pdf],out/'selected.pdf','2')
        office.pdf_edit('rotate',[pdf],out/'rotated.pdf','1',90)
        office.pdf_edit('merge',[pdf,pdf],out/'merged.pdf')
        results.append({'kind':'real_pdf_operations','status':'PASS','rendered':r['total_pages']})
    except Exception as e:results.append({'kind':'real_pdf_operations','status':'FAIL','error':str(e)})
    status='FAIL' if any(x['status']=='FAIL' for x in results) else ('PARTIAL' if any(x['status']=='NOT_RUN' for x in results) else 'PASS')
    result={'synthetic_test':True,'summary':status,'results':results,'visual_review':'NOT_PERFORMED',
            'notice':'Open the PNGs. This fixture cannot prove preservation of arbitrary customer files or Microsoft Office parity.'}
    new_file(out/'smoke-report.json',json_bytes(result))
    return result
